"""
File Path: /backend/app/plan/services/xlsx_io.py
Description: XLSX import/export helpers for Plan owner services.
Main Features:
    - Builds the Plan XLSX template and export workbook
    - Parses uploaded XLSX rows using the shared Plan field contract
    - Keeps file parsing separate from database writes
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from app.plan.domain import (
    PLAN_XLSX_COLUMNS,
    canonicalize_coil_no,
    decimal_cell_value,
    normalize_xlsx_header,
)
from app.plan.models import Plan
from app.plan.schemas.plan import PlanCreate, PlanXlsxImportError


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class ParsedPlanRow:
    row_number: int
    plan: PlanCreate
    canonical_coil_no: str


@dataclass(frozen=True)
class ParsedPlanWorkbook:
    rows: list[ParsedPlanRow]
    errors: list[PlanXlsxImportError]


_COLUMN_BY_ALIAS = {
    normalize_xlsx_header(alias): column
    for column in PLAN_XLSX_COLUMNS
    for alias in (column.header, column.key, *column.aliases)
}


def build_template_workbook() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "plans"
    _write_header(worksheet)
    worksheet.append([_xlsx_value(column.example) for column in PLAN_XLSX_COLUMNS])
    worksheet.freeze_panes = "A2"
    return _save_workbook(workbook)


def build_export_workbook(plans: list[Plan]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "plans"
    _write_header(worksheet)

    for plan in plans:
        worksheet.append([
            _xlsx_value(getattr(plan, column.key))
            for column in PLAN_XLSX_COLUMNS
        ])

    worksheet.freeze_panes = "A2"
    return _save_workbook(workbook)


def parse_plan_workbook(content: bytes) -> ParsedPlanWorkbook:
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        return ParsedPlanWorkbook(
            rows=[],
            errors=[
                PlanXlsxImportError(
                    row=None,
                    field=None,
                    error_code="plan_xlsx_invalid_file",
                    message="plan_xlsx_invalid_file",
                )
            ],
        )

    worksheet = workbook.active
    header_cells = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_cells:
        return ParsedPlanWorkbook(rows=[], errors=[_workbook_error("plan_xlsx_missing_header")])

    column_indexes: dict[int, str] = {}
    for index, value in enumerate(header_cells):
        normalized = normalize_xlsx_header(value)
        if not normalized:
            continue
        column = _COLUMN_BY_ALIAS.get(normalized)
        if column is not None:
            column_indexes[index] = column.key

    required_missing = [
        column.header
        for column in PLAN_XLSX_COLUMNS
        if column.required and column.key not in set(column_indexes.values())
    ]
    if required_missing:
        return ParsedPlanWorkbook(
            rows=[],
            errors=[
                _workbook_error(
                    "plan_xlsx_missing_required_columns",
                    field=",".join(required_missing),
                )
            ],
        )

    parsed_rows: list[ParsedPlanRow] = []
    errors: list[PlanXlsxImportError] = []
    seen_in_file: dict[str, int] = {}

    for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if _is_empty_row(cells):
            continue

        raw_data: dict[str, Any] = {}
        for index, field_key in column_indexes.items():
            if index >= len(cells):
                continue
            value = _clean_cell(cells[index])
            if value is not None:
                raw_data[field_key] = _typed_cell_value(field_key, value)

        for column in PLAN_XLSX_COLUMNS:
            if column.required and raw_data.get(column.key) is None:
                errors.append(
                    PlanXlsxImportError(
                        row=row_number,
                        field=column.key,
                        error_code="plan_xlsx_required_field_missing",
                        message="plan_xlsx_required_field_missing",
                    )
                )

        canonical_coil_no = canonicalize_coil_no(raw_data.get("coil_no"))
        if canonical_coil_no is None:
            errors.append(
                PlanXlsxImportError(
                    row=row_number,
                    field="coil_no",
                    error_code="plan_xlsx_invalid_coil_no",
                    message="plan_xlsx_invalid_coil_no",
                )
            )
        elif canonical_coil_no in seen_in_file:
            errors.append(
                PlanXlsxImportError(
                    row=row_number,
                    field="coil_no",
                    error_code="plan_xlsx_duplicate_coil_no_in_file",
                    message="plan_xlsx_duplicate_coil_no_in_file",
                    detail={"firstRow": seen_in_file[canonical_coil_no]},
                )
            )
        else:
            seen_in_file[canonical_coil_no] = row_number

        if any(error.row == row_number for error in errors):
            continue

        try:
            plan = PlanCreate.model_validate(raw_data)
        except ValidationError as exc:
            errors.extend(_validation_errors(row_number, exc))
            continue

        parsed_rows.append(
            ParsedPlanRow(
                row_number=row_number,
                plan=plan,
                canonical_coil_no=canonical_coil_no or "",
            )
        )

    return ParsedPlanWorkbook(rows=parsed_rows, errors=errors)


def _write_header(worksheet) -> None:
    worksheet.append([column.header for column in PLAN_XLSX_COLUMNS])
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for index, column in enumerate(PLAN_XLSX_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        worksheet.column_dimensions[cell.column_letter].width = max(14, len(column.header) + 2)


def _save_workbook(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _clean_cell(value: object) -> object | None:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return value


def _typed_cell_value(field_key: str, value: object) -> object:
    if field_key == "priority":
        decimal_value = decimal_cell_value(value)
        return int(decimal_value) if decimal_value is not None else value
    if field_key in _decimal_fields():
        return decimal_cell_value(value) if decimal_cell_value(value) is not None else value
    return value


def _decimal_fields() -> set[str]:
    return {
        "weight",
        "thickness",
        "width",
        "density",
        "entry_length",
        "uncoiler1_tension",
        "exit_tension",
        "process1_tension",
        "entry_loop_tension",
        "exit_loop_tension",
        "entry_unit_tension",
        "entry_loop_unit_tension",
        "process_unit_tension",
        "exit_loop_unit_tension",
        "exit_unit_tension",
    }


def _is_empty_row(cells: tuple[object, ...]) -> bool:
    return all(_clean_cell(value) is None for value in cells)


def _workbook_error(error_code: str, *, field: str | None = None) -> PlanXlsxImportError:
    return PlanXlsxImportError(row=None, field=field, error_code=error_code, message=error_code)


def _validation_errors(row_number: int, exc: ValidationError) -> list[PlanXlsxImportError]:
    errors: list[PlanXlsxImportError] = []
    for error in exc.errors():
        field = ".".join(str(item) for item in error["loc"]) if error.get("loc") else None
        errors.append(
            PlanXlsxImportError(
                row=row_number,
                field=field,
                error_code="plan_xlsx_validation_failed",
                message="plan_xlsx_validation_failed",
                detail={"type": error.get("type")},
            )
        )
    return errors


def _xlsx_value(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return value
