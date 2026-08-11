"""
File Path: /tools/lineage-mapping-studio/src-python/graph/workbook.py
Description: Projection workbook reader for Field -> PLC graph bindings.
Main Features:
  - Reads PLC point reference rows from PLC_POINTS.
  - Reads table field to PLC point bindings from business table sheets.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from .models import ToolError, value


SKIP_SHEETS = {"README", "PLC_POINTS", "TABLE_INDEX"}


def read_workbook_bindings(path: Path | None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if path is None or not path.exists():
        return [], []
    try:
        import openpyxl
    except ImportError as exc:
        raise ToolError("openpyxl is required to read plc_projection_mapping.xlsx") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    workbook_points: list[dict[str, Any]] = []
    bindings: list[dict[str, Any]] = []

    if "PLC_POINTS" in workbook.sheetnames:
        workbook_points = sheet_rows(workbook["PLC_POINTS"])

    for sheet_name in workbook.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        for row in sheet_rows(workbook[sheet_name]):
            field_name = value(row.get("field_name"))
            point_name = value(row.get("point_name"))
            source_type = value(row.get("source_type"))
            if not field_name:
                continue
            if not point_name and source_type != "plc":
                continue
            bindings.append(
                {
                    "table_name": sheet_name,
                    "field_name": field_name,
                    "field_type": value(row.get("field_type")),
                    "source_type": source_type,
                    "snapshot_source": value(row.get("snapshot_source")),
                    "projection_mode": value(row.get("projection_mode")),
                    "plc_key": value(row.get("plc_key")),
                    "db_number": row.get("db_number"),
                    "group_name": value(row.get("group_name")),
                    "point_name": point_name,
                    "source_name": value(row.get("source_name")),
                    "plc_data_type": value(row.get("plc_data_type")),
                    "projection_policy": value(row.get("projection_policy")),
                    "transform_rule": value(row.get("transform_rule")),
                    "notes": value(row.get("notes")),
                }
            )
    return bindings, workbook_points


def sheet_rows(worksheet: Any) -> list[dict[str, Any]]:
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = [value(item) for item in next(rows)]
    except StopIteration:
        return []
    result: list[dict[str, Any]] = []
    for values in rows:
        item = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        if any(value(cell) for cell in item.values()):
            result.append(item)
    return result
