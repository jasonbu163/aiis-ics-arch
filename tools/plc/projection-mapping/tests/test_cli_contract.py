"""
File Path: /tools/plc/projection-mapping/tests/test_cli_contract.py
Description: CLI contract tests for projection-mapping.
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path
import unittest

from openpyxl import load_workbook


TOOL_ROOT = Path(__file__).resolve().parents[1]
MAIN_PY = TOOL_ROOT / "main.py"


class CliContractTests(unittest.TestCase):
    def test_bare_entry_prints_help_without_action(self) -> None:
        result = self._run()

        self.assertEqual(result.returncode, 0)
        self.assertIn("PLC projection mapping helper", result.stdout)
        self.assertIn("generate-xlsx", result.stdout)

    def test_unknown_command_fails_with_suggestion(self) -> None:
        result = self._run("generate-xlxs")

        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Did you mean 'generate-xlsx'", result.stderr)

    def test_flag_without_explicit_command_is_rejected(self) -> None:
        result = self._run("--snapshot-policy", "tools/plc/projection-mapping/inputs/plc_snapshot_policy.yaml")

        self.assertNotEqual(result.returncode, 0)
        self.assertIn("explicit command required", result.stderr)

    def test_old_snapshot_policy_shape_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            targets_path = temp_path / "projection_targets.json"
            schema_path = temp_path / "table_schema.json"
            policy_path = temp_path / "plc_snapshot_policy.yaml"
            targets_path.write_text(json.dumps({"tables": [{"table": "monitor"}]}), encoding="utf-8")
            schema_path.write_text(
                json.dumps({
                    "tables": [
                        {
                            "table_name": "monitor",
                            "columns": [
                                {"field_name": "temperature", "field_type": "float", "nullable": True},
                            ],
                        }
                    ]
                }),
                encoding="utf-8",
            )
            policy_path.write_text(
                """\
version: 2
groups: []
""",
                encoding="utf-8",
            )

            result = self._run(
                "validate-inputs",
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
            )

        self.assertNotEqual(result.returncode, 0)
        self.assertIn("old top-level version/groups shape", result.stdout)

    def test_generate_xlsx_requires_overwrite_for_existing_output(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            targets_path = temp_path / "projection_targets.json"
            schema_path = temp_path / "table_schema.json"
            policy_path = temp_path / "plc_snapshot_policy.yaml"
            output_path = temp_path / "plc_projection_mapping.xlsx"
            targets_path.write_text(json.dumps({"tables": [{"table": "monitor"}]}), encoding="utf-8")
            schema_path.write_text(
                json.dumps({
                    "tables": [
                        {
                            "table_name": "monitor",
                            "columns": [
                                {"field_name": "temperature", "field_type": "float", "nullable": True},
                            ],
                        }
                    ]
                }),
                encoding="utf-8",
            )
            policy_path.write_text(self._minimal_policy(), encoding="utf-8")

            first = self._run(
                "generate-xlsx",
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
                "--output",
                str(output_path),
            )
            second = self._run(
                "generate-xlsx",
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
                "--output",
                str(output_path),
            )
            third = self._run(
                "generate-xlsx",
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
                "--output",
                str(output_path),
                "--overwrite",
            )

        self.assertEqual(first.returncode, 0)
        self.assertNotEqual(second.returncode, 0)
        self.assertIn("pass --overwrite", second.stdout)
        self.assertEqual(third.returncode, 0)

    def test_legacy_combined_projection_mode_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            targets_path = temp_path / "projection_targets.json"
            schema_path = temp_path / "table_schema.json"
            policy_path = temp_path / "plc_snapshot_policy.yaml"
            output_path = temp_path / "plc_projection_mapping.xlsx"
            targets_path.write_text(json.dumps({"tables": [{"table": "monitor"}]}), encoding="utf-8")
            schema_path.write_text(
                json.dumps({
                    "tables": [
                        {
                            "table_name": "monitor",
                            "columns": [
                                {"field_name": "temperature", "field_type": "float", "nullable": True},
                            ],
                        }
                    ]
                }),
                encoding="utf-8",
            )
            policy_path.write_text(self._minimal_policy(), encoding="utf-8")

            created = self._run(
                "generate-xlsx",
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
                "--output",
                str(output_path),
            )
            self.assertEqual(created.returncode, 0)

            workbook = load_workbook(output_path)
            workbook["TABLE_INDEX"]["D2"] = "raw/append_history"
            workbook.save(output_path)

            result = self._run(
                "validate-xlsx",
                "--mapping",
                str(output_path),
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
            )

        self.assertNotEqual(result.returncode, 0)
        self.assertIn("invalid projection_mode 'raw/append_history'", result.stdout)

    def test_generate_xlsx_uses_snapshot_policy_candidates_only(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            targets_path = temp_path / "projection_targets.json"
            schema_path = temp_path / "table_schema.json"
            policy_path = temp_path / "plc_snapshot_policy.yaml"
            output_path = temp_path / "plc_projection_mapping.xlsx"
            targets_path.write_text(json.dumps({"tables": [{"table": "monitor"}]}), encoding="utf-8")
            schema_path.write_text(
                json.dumps({
                    "tables": [
                        {
                            "table_name": "monitor",
                            "columns": [
                                {"field_name": "temperature", "field_type": "float", "nullable": True},
                            ],
                        }
                    ]
                }),
                encoding="utf-8",
            )
            policy_path.write_text(self._policy_with_disabled_point(), encoding="utf-8")

            result = self._run(
                "generate-xlsx",
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
                "--output",
                str(output_path),
            )

            workbook = load_workbook(output_path, read_only=True)
            rows = list(workbook["PLC_POINTS"].iter_rows(values_only=True))

        self.assertEqual(result.returncode, 0)
        self.assertIn('"policy_point_count": 3', result.stdout)
        self.assertIn('"candidate_point_count": 2', result.stdout)
        self.assertEqual(rows[0], (
            "plc_key",
            "db_number",
            "group_name",
            "point_name",
            "source_name",
            "plc_data_type",
            "offset",
            "bit",
            "desc",
            "raw_enabled",
            "raw_policy",
            "latest_enabled",
            "available_sources",
        ))
        point_names = [row[3] for row in rows[1:]]
        self.assertEqual(point_names, ["DB28_RawOnly_Real", "DB28_LatestOnly_Real"])
        self.assertEqual(rows[1][12], "raw")
        self.assertEqual(rows[2][12], "latest")

    def test_validate_xlsx_rejects_source_disabled_by_policy(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            targets_path = temp_path / "projection_targets.json"
            schema_path = temp_path / "table_schema.json"
            policy_path = temp_path / "plc_snapshot_policy.yaml"
            output_path = temp_path / "plc_projection_mapping.xlsx"
            targets_path.write_text(json.dumps({"tables": [{"table": "monitor"}]}), encoding="utf-8")
            schema_path.write_text(
                json.dumps({
                    "tables": [
                        {
                            "table_name": "monitor",
                            "columns": [
                                {"field_name": "temperature", "field_type": "float", "nullable": True},
                            ],
                        }
                    ]
                }),
                encoding="utf-8",
            )
            policy_path.write_text(self._policy_with_disabled_point(), encoding="utf-8")

            created = self._run(
                "generate-xlsx",
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
                "--output",
                str(output_path),
            )
            self.assertEqual(created.returncode, 0)

            workbook = load_workbook(output_path)
            sheet = workbook["monitor"]
            sheet["E2"] = "plc"
            sheet["F2"] = "latest"
            sheet["G2"] = "direct_read"
            sheet["H2"] = "PLC_1"
            sheet["I2"] = 28
            sheet["J2"] = "Example_Read_Data"
            sheet["K2"] = "DB28_RawOnly_Real"
            sheet["L2"] = "Real"
            workbook.save(output_path)

            result = self._run(
                "validate-xlsx",
                "--mapping",
                str(output_path),
                "--targets",
                str(targets_path),
                "--schema",
                str(schema_path),
                "--snapshot-policy",
                str(policy_path),
            )

        self.assertNotEqual(result.returncode, 0)
        self.assertIn("requires latest_enabled=true", result.stdout)

    def _run(self, *args: str) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            [sys.executable, str(MAIN_PY), *args],
            cwd=TOOL_ROOT.parents[2],
            text=True,
            capture_output=True,
            check=False,
        )

    def _minimal_policy(self) -> str:
        return """\
PLC_1:
  snapshot_policy_version: 2
  point_count: 1
  raw_point_count: 1
  latest_point_count: 1
  groups:
  - name: Example_Read_Data
    db_number: 28
    point_count: 1
    raw_point_count: 1
    latest_point_count: 1
    points:
    - name: DB28_ActTemp_Z1_Real
      source_name: ActTemp_Z1
      type: Real
      offset: 20
      desc: Temperature
      raw_enabled: true
      raw_policy: every_sample
      latest_enabled: true
      reason: ""
"""

    def _policy_with_disabled_point(self) -> str:
        return """\
PLC_1:
  snapshot_policy_version: 2
  point_count: 3
  raw_point_count: 1
  latest_point_count: 1
  groups:
  - name: Example_Read_Data
    db_number: 28
    point_count: 3
    raw_point_count: 1
    latest_point_count: 1
    points:
    - name: DB28_RawOnly_Real
      source_name: RawOnly
      type: Real
      offset: 20
      desc: Raw only
      raw_enabled: true
      raw_policy: every_sample
      latest_enabled: false
      reason: ""
    - name: DB28_LatestOnly_Real
      source_name: LatestOnly
      type: Real
      offset: 24
      desc: Latest only
      raw_enabled: false
      raw_policy: every_sample
      latest_enabled: true
      reason: ""
    - name: DB28_Disabled_Real
      source_name: Disabled
      type: Real
      offset: 28
      desc: Disabled
      raw_enabled: false
      raw_policy: every_sample
      latest_enabled: false
      reason: ""
"""


if __name__ == "__main__":
    unittest.main()
