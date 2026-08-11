"""
File Path: /tools/plc/projection-mapping/src-python/core/workbook.py
Description: Projection mapping workbook generation.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from core.common import ToolError, write_yaml
from core.points import validate_inputs


BUSINESS_SHEET_HEADERS = [
    "field_name",
    "field_type",
    "nullable",
    "business_desc",
    "source_type",
    "snapshot_source",
    "projection_mode",
    "plc_key",
    "db_number",
    "group_name",
    "point_name",
    "plc_data_type",
    "transform_rule",
    "default_value",
    "unit",
    "is_required",
    "projection_policy",
    "unique_key_group",
    "notes",
]

TABLE_INDEX_HEADERS = [
    "table_name",
    "field_count",
    "snapshot_source",
    "projection_mode",
    "projection_policy",
    "unique_key_group",
    "purpose",
    "notes",
]

INHERITABLE_COLUMNS = [
    "snapshot_source",
    "projection_mode",
    "projection_policy",
    "unique_key_group",
]

MAPPING_SIGNAL_COLUMNS = [
    "source_type",
    "snapshot_source",
    "projection_mode",
    "plc_key",
    "db_number",
    "group_name",
    "point_name",
    "plc_data_type",
    "transform_rule",
    "unit",
    "is_required",
    "projection_policy",
    "unique_key_group",
]

VALID_SOURCE_TYPES = {"plc", "system", "manual", "calculated", "constant"}
VALID_SNAPSHOT_SOURCES = {"latest", "raw"}
VALID_PROJECTION_MODES = {"direct_read", "upsert_current", "append_history", "state_machine"}
VALID_PROJECTION_POLICIES = {"latest_only", "every_sample", "on_change", "hourly"}
TRUE_VALUES = {"true", "1", "yes", "y"}
FALSE_VALUES = {"false", "0", "no", "n"}


def add_sheet_rows(sheet: Any, rows: list[list[Any]]) -> None:
    for row in rows:
        sheet.append(row)


def autosize_columns(sheet: Any, max_width: int = 48) -> None:
    for column_cells in sheet.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if not values:
            continue
        width = min(max(len(value) for value in values) + 2, max_width)
        sheet.column_dimensions[column_cells[0].column_letter].width = width


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def normalized_token(value: Any) -> str:
    return cell_text(value).lower()


def parse_optional_bool(value: Any) -> bool | None:
    token = normalized_token(value)
    if not token:
        return None
    if token in TRUE_VALUES:
        return True
    if token in FALSE_VALUES:
        return False
    raise ValueError(cell_text(value))


def sheet_headers(sheet: Any) -> dict[str, int]:
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {
        cell_text(value): index
        for index, value in enumerate(header_row)
        if cell_text(value)
    }


def row_dict(headers: dict[str, int], row: tuple[Any, ...]) -> dict[str, Any]:
    return {
        header: row[index] if index < len(row) else None
        for header, index in headers.items()
    }


def sheet_name_for_table(table_name: str) -> str:
    return table_name[:31]


def has_mapping_signal(row: dict[str, Any]) -> bool:
    return any(cell_text(row.get(column)) for column in MAPPING_SIGNAL_COLUMNS)


def clean_mapping(row: dict[str, Any], columns: list[str]) -> dict[str, str]:
    return {
        column: cell_text(row.get(column))
        for column in columns
        if cell_text(row.get(column))
    }


def read_table_defaults(workbook: Any, targets: list[str], errors: list[str]) -> dict[str, dict[str, str]]:
    defaults = {table_name: {column: "" for column in INHERITABLE_COLUMNS} for table_name in targets}
    if "TABLE_INDEX" not in workbook.sheetnames:
        errors.append("missing TABLE_INDEX sheet")
        return defaults

    sheet = workbook["TABLE_INDEX"]
    headers = sheet_headers(sheet)
    if "table_name" not in headers:
        errors.append("TABLE_INDEX missing required table_name header")
        return defaults

    target_set = set(targets)
    for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = row_dict(headers, row)
        table_name = cell_text(values.get("table_name"))
        if not table_name:
            continue
        if table_name not in target_set:
            continue
        for column in INHERITABLE_COLUMNS:
            value = cell_text(values.get(column))
            if column in {"snapshot_source", "projection_mode", "projection_policy"}:
                value = value.lower()
            defaults[table_name][column] = value

        snapshot_source = defaults[table_name]["snapshot_source"]
        projection_mode = defaults[table_name]["projection_mode"]
        projection_policy = defaults[table_name]["projection_policy"]
        if snapshot_source and snapshot_source not in VALID_SNAPSHOT_SOURCES:
            errors.append(f"TABLE_INDEX row {excel_row}: invalid snapshot_source {snapshot_source!r}")
        if projection_mode and projection_mode not in VALID_PROJECTION_MODES:
            errors.append(f"TABLE_INDEX row {excel_row}: invalid projection_mode {projection_mode!r}")
        if projection_policy and projection_policy not in VALID_PROJECTION_POLICIES:
            errors.append(f"TABLE_INDEX row {excel_row}: invalid projection_policy {projection_policy!r}")
    return defaults


def effective_value(row: dict[str, Any], defaults: dict[str, str], column: str, source_type: str) -> str:
    value = cell_text(row.get(column))
    if column in {"snapshot_source", "projection_mode", "projection_policy"}:
        value = value.lower()
    if value:
        return value
    if source_type == "plc":
        return defaults.get(column, "")
    return ""


def build_field_rule(
    *,
    table_name: str,
    excel_row: int,
    row: dict[str, Any],
    table_defaults: dict[str, str],
    point_by_name: dict[str, dict[str, Any]],
    errors: list[str],
) -> dict[str, Any] | None:
    if not has_mapping_signal(row):
        return None

    field_name = cell_text(row.get("field_name"))
    source_type = normalized_token(row.get("source_type"))
    location = f"{table_name} row {excel_row}"
    if not field_name:
        errors.append(f"{location}: missing field_name")
        return None
    if not source_type:
        errors.append(f"{location} field {field_name}: source_type is required when mapping columns are filled")
        return None
    if source_type not in VALID_SOURCE_TYPES:
        errors.append(f"{location} field {field_name}: invalid source_type {source_type!r}")
        return None

    try:
        is_required = parse_optional_bool(row.get("is_required"))
    except ValueError as exc:
        errors.append(f"{location} field {field_name}: invalid is_required {exc.args[0]!r}")
        is_required = None

    snapshot_source = effective_value(row, table_defaults, "snapshot_source", source_type)
    projection_mode = effective_value(row, table_defaults, "projection_mode", source_type)
    projection_policy = effective_value(row, table_defaults, "projection_policy", source_type)
    unique_key_group = effective_value(row, table_defaults, "unique_key_group", source_type)

    if snapshot_source and snapshot_source not in VALID_SNAPSHOT_SOURCES:
        errors.append(f"{location} field {field_name}: invalid snapshot_source {snapshot_source!r}")
    if projection_mode and projection_mode not in VALID_PROJECTION_MODES:
        errors.append(f"{location} field {field_name}: invalid projection_mode {projection_mode!r}")
    if projection_policy and projection_policy not in VALID_PROJECTION_POLICIES:
        errors.append(f"{location} field {field_name}: invalid projection_policy {projection_policy!r}")

    rule: dict[str, Any] = {
        "table": table_name,
        "field": field_name,
        "field_type": cell_text(row.get("field_type")),
        "source_type": source_type,
    }
    for source_column, rule_column in [
        ("business_desc", "business_desc"),
        ("transform_rule", "transform_rule"),
        ("default_value", "default_value"),
        ("unit", "unit"),
        ("notes", "notes"),
    ]:
        value = cell_text(row.get(source_column))
        if value:
            rule[rule_column] = value
    if is_required is not None:
        rule["is_required"] = is_required
    if snapshot_source:
        rule["snapshot_source"] = snapshot_source
    if projection_mode:
        rule["projection_mode"] = projection_mode
    if projection_policy:
        rule["projection_policy"] = projection_policy
    if unique_key_group:
        rule["unique_key_group"] = unique_key_group

    if source_type != "plc":
        return rule

    if not snapshot_source:
        errors.append(f"{location} field {field_name}: snapshot_source is required for PLC fields")
    if not projection_mode:
        errors.append(f"{location} field {field_name}: projection_mode is required for PLC fields")

    point_name = cell_text(row.get("point_name"))
    if not point_name:
        errors.append(f"{location} field {field_name}: point_name is required for PLC fields")
        return rule

    point = point_by_name.get(point_name)
    if not point:
        errors.append(f"{location} field {field_name}: unknown point_name {point_name!r}")
        return rule
    if snapshot_source == "latest" and not point.get("latest_enabled"):
        errors.append(f"{location} field {field_name}: point_name {point_name!r} requires latest_enabled=true")
    if snapshot_source == "raw" and not point.get("raw_enabled"):
        errors.append(f"{location} field {field_name}: point_name {point_name!r} requires raw_enabled=true")

    for column in ["plc_key", "db_number", "group_name", "plc_data_type"]:
        row_value = cell_text(row.get(column))
        point_key = "plc_data_type" if column == "plc_data_type" else column
        point_value = cell_text(point.get(point_key))
        if row_value and row_value.lower() != point_value.lower():
            errors.append(
                f"{location} field {field_name}: {column} {row_value!r} does not match PLC_POINTS value {point_value!r}"
            )

    plc_payload = {
        "plc_key": point.get("plc_key"),
        "db_number": point.get("db_number"),
        "group_name": point.get("group_name"),
        "point_name": point_name,
        "plc_data_type": point.get("plc_data_type"),
        "offset": point.get("offset"),
        "bit": point.get("bit"),
    }
    rule["plc"] = {
        key: value
        for key, value in plc_payload.items()
        if value is not None and cell_text(value)
    }
    return rule


def load_mapping_workbook(
    mapping_path: Path,
    targets_path: Path,
    schema_path: Path,
    snapshot_policy_path: Path,
) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ToolError("openpyxl is required to read mapping workbooks") from exc

    state = validate_inputs(targets_path, schema_path, snapshot_policy_path)
    targets: list[str] = state["targets"]
    schema: dict[str, list[dict[str, Any]]] = state["schema"]
    point_by_name = {point["point_name"]: point for point in state["points"]}

    try:
        workbook = load_workbook(mapping_path, data_only=True, read_only=True)
    except FileNotFoundError as exc:
        raise ToolError(f"missing XLSX file: {mapping_path}") from exc
    except Exception as exc:
        raise ToolError(f"cannot read XLSX file: {mapping_path}: {exc}") from exc

    errors: list[str] = []
    table_defaults = read_table_defaults(workbook, targets, errors)
    rules: list[dict[str, Any]] = []
    total_field_rows = 0
    skipped_field_rows = 0
    inherited_value_count = 0

    for table_name in targets:
        sheet_name = sheet_name_for_table(table_name)
        if sheet_name not in workbook.sheetnames:
            errors.append(f"missing business table sheet {sheet_name!r} for table {table_name!r}")
            continue
        sheet = workbook[sheet_name]
        headers = sheet_headers(sheet)
        missing_headers = [header for header in ["field_name", "source_type", "point_name"] if header not in headers]
        if missing_headers:
            errors.append(f"{table_name}: missing headers {missing_headers}")
            continue

        schema_fields = {
            cell_text(column.get("field_name") or column.get("name"))
            for column in schema[table_name]
        }
        seen_fields: set[str] = set()
        for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = row_dict(headers, row)
            field_name = cell_text(values.get("field_name"))
            if not field_name:
                continue
            total_field_rows += 1
            if field_name in seen_fields:
                errors.append(f"{table_name} row {excel_row}: duplicate field_name {field_name!r}")
            seen_fields.add(field_name)
            if field_name not in schema_fields:
                errors.append(f"{table_name} row {excel_row}: field_name {field_name!r} not found in table_schema.json")

            before_rule_errors = len(errors)
            rule = build_field_rule(
                table_name=table_name,
                excel_row=excel_row,
                row=values,
                table_defaults=table_defaults[table_name],
                point_by_name=point_by_name,
                errors=errors,
            )
            if rule is None:
                skipped_field_rows += 1
                continue
            if len(errors) == before_rule_errors:
                for column in INHERITABLE_COLUMNS:
                    if column in rule and not cell_text(values.get(column)) and table_defaults[table_name].get(column):
                        inherited_value_count += 1
                rules.append(rule)

    if errors:
        examples = "; ".join(errors[:20])
        if len(errors) > 20:
            examples = f"{examples}; ... ({len(errors)} total errors)"
        raise ToolError(f"invalid mapping workbook: {examples}")

    return {
        "targets": targets,
        "schema": schema,
        "table_defaults": table_defaults,
        "rules": rules,
        "summary": {
            "workbook": str(mapping_path),
            "table_count": len(targets),
            "field_row_count": total_field_rows,
            "mapped_row_count": len(rules),
            "plc_rule_count": sum(1 for rule in rules if rule["source_type"] == "plc"),
            "skipped_row_count": skipped_field_rows,
            "inherited_value_count": inherited_value_count,
            "policy_point_count": state["policy_point_count"],
            "candidate_point_count": state["candidate_point_count"],
            "raw_enabled_count": state["raw_enabled_count"],
            "latest_enabled_count": state["latest_enabled_count"],
        },
    }


def validate_mapping_workbook(
    mapping_path: Path,
    targets_path: Path,
    schema_path: Path,
    snapshot_policy_path: Path,
) -> dict[str, Any]:
    return load_mapping_workbook(mapping_path, targets_path, schema_path, snapshot_policy_path)["summary"]


def generate_rules(
    mapping_path: Path,
    targets_path: Path,
    schema_path: Path,
    snapshot_policy_path: Path,
    output_path: Path,
) -> dict[str, Any]:
    parsed = load_mapping_workbook(mapping_path, targets_path, schema_path, snapshot_policy_path)
    rules_by_table: dict[str, list[dict[str, Any]]] = {table_name: [] for table_name in parsed["targets"]}
    for rule in parsed["rules"]:
        field_rule = {key: value for key, value in rule.items() if key != "table"}
        rules_by_table[rule["table"]].append(field_rule)

    payload = {
        "version": 1,
        "source": {
            "workbook": str(mapping_path),
            "target_count": len(parsed["targets"]),
            "rule_count": len(parsed["rules"]),
        },
        "tables": [
            {
                "table": table_name,
                "defaults": {
                    key: value
                    for key, value in parsed["table_defaults"][table_name].items()
                    if value
                },
                "fields": rules_by_table[table_name],
            }
            for table_name in parsed["targets"]
            if rules_by_table[table_name] or any(parsed["table_defaults"][table_name].values())
        ],
    }
    write_yaml(output_path, payload)
    return {
        "output": str(output_path),
        **parsed["summary"],
    }


def generate_xlsx(targets_path: Path, schema_path: Path, snapshot_policy_path: Path, output_path: Path) -> dict[str, Any]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ToolError("openpyxl is required for generate-xlsx") from exc

    state = validate_inputs(targets_path, schema_path, snapshot_policy_path)
    targets: list[str] = state["targets"]
    schema: dict[str, list[dict[str, Any]]] = state["schema"]
    points: list[dict[str, Any]] = state["points"]

    workbook = Workbook()
    workbook.remove(workbook.active)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    title_font = Font(bold=True)

    readme = workbook.create_sheet("README")
    add_sheet_rows(readme, [
        ["Purpose", "Fill PLC point bindings for selected business table fields."],
        ["Flow", "export-schema -> review table_schema.json -> generate-xlsx -> fill workbook -> generate-rules."],
        ["Do not edit", "PLC_POINTS is generated from plc_snapshot_policy.yaml and should be treated as read-only reference."],
        ["Required fields", "source_type is required for mapped rows. PLC rows also require point_name plus effective snapshot_source and projection_mode."],
        ["TABLE_INDEX defaults", "PLC field rows may inherit snapshot_source, projection_mode, projection_policy and unique_key_group from TABLE_INDEX."],
        ["snapshot_source", "latest for direct/current reads; raw for history, reports, state machines and cursor processing."],
        ["projection_mode", "direct_read, upsert_current, append_history or state_machine."],
    ])
    autosize_columns(readme)

    point_sheet = workbook.create_sheet("PLC_POINTS")
    point_headers = [
        "plc_key",
        "db_number",
        "group_name",
        "point_name",
        "source_name",
        "plc_data_type",
        "offset",
        "bit",
        "desc",
        "raw_enabled",
        "raw_policy",
        "latest_enabled",
        "available_sources",
    ]
    point_sheet.append(point_headers)
    for point in points:
        point_sheet.append([point.get(header) for header in point_headers])

    index_sheet = workbook.create_sheet("TABLE_INDEX")
    index_sheet.append(TABLE_INDEX_HEADERS)
    for table_name in targets:
        index_sheet.append([table_name, len(schema[table_name]), "", "", "", "", "", ""])

    for table_name in targets:
        sheet_name = table_name[:31]
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(BUSINESS_SHEET_HEADERS)
        for column in schema[table_name]:
            field_name = column.get("field_name") or column.get("name")
            field_type = column.get("field_type") or column.get("data_type")
            nullable = column.get("nullable")
            comment = column.get("comment") or ""
            sheet.append([
                field_name,
                field_type,
                nullable,
                comment,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                column.get("default"),
                "",
                "",
                "",
                "",
                "",
            ])
        for cell in sheet[1]:
            cell.font = title_font
            cell.fill = header_fill
        sheet.freeze_panes = "A2"
        for index, _header in enumerate(BUSINESS_SHEET_HEADERS, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = 18

    for sheet in workbook.worksheets:
        if sheet.max_row:
            for cell in sheet[1]:
                cell.font = title_font
                cell.fill = header_fill
        autosize_columns(sheet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {
        "output": str(output_path),
        "target_count": len(targets),
        "point_count": len(points),
        "policy_point_count": state["policy_point_count"],
        "candidate_point_count": state["candidate_point_count"],
        "raw_enabled_count": state["raw_enabled_count"],
        "latest_enabled_count": state["latest_enabled_count"],
        "table_count": len(targets),
    }
