"""
文件路径: /tools/plc/snapshot-policy/src-python/core/workbook.py
功能描述: PLC 快照 raw/latest 范围策略 XLSX 模板生成与读取校验
"""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core.models import POLICY_COLUMNS, POLICY_OPTIONS, PlcPoint, SnapshotPolicyRow


CHINESE_PATTERN = re.compile(r"[\u4e00-\u9fff]")


def contains_chinese(*values: str) -> bool:
    return any(CHINESE_PATTERN.search(value or "") for value in values)


def default_raw_enabled(point: PlcPoint) -> bool:
    return contains_chinese(point.source_name, point.desc)


def default_raw_policy(point: PlcPoint) -> str:
    return "every_sample"


def default_latest_enabled(point: PlcPoint) -> bool:
    return True


def default_reason(point: PlcPoint) -> str:
    if default_raw_enabled(point):
        return "source_name_or_desc_contains_chinese"
    return "manual_review_required"


def parse_bool(value: Any, *, default: bool | None = None) -> bool | None:
    if isinstance(value, bool):
        return value
    if value is None or str(value).strip() == "":
        return default
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "y", "是", "启用"}:
        return True
    if text in {"false", "0", "no", "n", "否", "禁用"}:
        return False
    return default


def required_bool(value: Any, field_name: str, row_label: str) -> bool:
    parsed = parse_bool(value)
    if parsed is None:
        raise ValueError(f"{row_label}: {field_name} must be TRUE or FALSE")
    return parsed


def required_int(value: Any, field_name: str, row_label: str) -> int:
    if value is None or str(value).strip() == "":
        raise ValueError(f"{row_label}: {field_name} is required")
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{row_label}: {field_name} must be an integer") from exc


def write_policy_workbook(points: list[PlcPoint], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["PLC snapshot raw/latest point-scope policy"])
    readme.append(["1. This file is generated from tools/config/plc_points.yaml."])
    readme.append(["2. Edit raw_enabled / raw_policy / latest_enabled / reason, then copy it to tools/plc/snapshot-policy/inputs/plc_snapshot_policy.xlsx."])
    readme.append(["3. One sheet is generated per DB/group so duplicate point names remain easy to review."])
    readme.append(["4. Run generate-yaml to read the inputs workbook and write outputs/plc_snapshot_policy.yaml."])
    readme.append(["5. raw_enabled controls raw decoded payload scope; latest_enabled controls latest decoded payload scope."])
    readme.append(["6. Default raw_enabled is TRUE when source_name or desc contains Chinese, otherwise FALSE."])
    readme.append(["7. Default latest_enabled is TRUE so operators explicitly review reductions before CA shrinks latest payloads."])
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 120

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    points_by_group: dict[tuple[int, str], list[PlcPoint]] = defaultdict(list)
    for point in points:
        points_by_group[(point.db_number, point.group_name)].append(point)

    used_sheet_names: set[str] = {"README"}
    for (db_number, group_name), group_points in sorted(points_by_group.items()):
        sheet = workbook.create_sheet(title=safe_sheet_name(db_number, group_name, used_sheet_names))
        raw_enabled_validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
        policy_validation = DataValidation(type="list", formula1=f'"{",".join(POLICY_OPTIONS)}"', allow_blank=False)
        latest_enabled_validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
        sheet.add_data_validation(raw_enabled_validation)
        sheet.add_data_validation(policy_validation)
        sheet.add_data_validation(latest_enabled_validation)
        sheet.append(POLICY_COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        for point in sorted(group_points, key=lambda item: (item.offset, item.bit or 0, item.name)):
            sheet.append([
                point.db_number,
                point.group_name,
                point.name,
                point.source_name,
                point.data_type,
                point.offset,
                point.bit if point.bit is not None else "",
                point.desc,
                default_raw_enabled(point),
                default_raw_policy(point),
                default_latest_enabled(point),
                default_reason(point),
            ])

        format_policy_sheet(sheet, raw_enabled_validation, policy_validation, latest_enabled_validation)

    workbook.save(output_path)


def read_policy_workbook(points: list[PlcPoint], workbook_path: Path) -> list[SnapshotPolicyRow]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Snapshot policy workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, data_only=True)
    point_by_key = {point_key(point): point for point in points}
    sheets = policy_sheets(workbook)
    if not sheets:
        raise ValueError("Workbook must contain at least one policy sheet")

    rows: list[SnapshotPolicyRow] = []
    seen_points: set[str] = set()
    for sheet in sheets:
        rows.extend(read_policy_sheet(sheet, point_by_key, seen_points))

    return rows


def read_policy_sheet(
    sheet: Any,
    point_by_key: dict[tuple[int, str, str], PlcPoint],
    seen_points: set[str],
) -> list[SnapshotPolicyRow]:
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    missing = [column for column in POLICY_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"{sheet.title} missing columns: {', '.join(missing)}")
    column_index = {name: headers.index(name) + 1 for name in POLICY_COLUMNS}

    rows: list[SnapshotPolicyRow] = []
    for row_number in range(2, sheet.max_row + 1):
        raw_name = sheet.cell(row=row_number, column=column_index["name"]).value
        if raw_name is None or str(raw_name).strip() == "":
            continue
        name = str(raw_name).strip()
        row_label = f"{sheet.title} row {row_number} ({name})"
        db_number = required_int(sheet.cell(row=row_number, column=column_index["db_number"]).value, "db_number", row_label)
        offset = required_int(sheet.cell(row=row_number, column=column_index["offset"]).value, "offset", row_label)
        data_type = str(sheet.cell(row=row_number, column=column_index["type"]).value or "").strip()
        group_name = str(sheet.cell(row=row_number, column=column_index["group_name"]).value or "").strip()
        point = point_by_key.get((db_number, group_name, name))
        if point is None:
            raise ValueError(f"{row_label}: point is not defined in plc_points.yaml")
        if db_number != point.db_number or offset != point.offset or data_type != point.data_type or group_name != point.group_name:
            raise ValueError(f"{row_label}: db_number/group_name/type/offset does not match plc_points.yaml")

        point_identity = f"DB{db_number}:{group_name}:{name}"
        if point_identity in seen_points:
            raise ValueError(f"{row_label}: duplicate policy point")
        seen_points.add(point_identity)

        raw_enabled = required_bool(sheet.cell(row=row_number, column=column_index["raw_enabled"]).value, "raw_enabled", row_label)
        raw_policy = str(sheet.cell(row=row_number, column=column_index["raw_policy"]).value or "").strip()
        if raw_policy not in POLICY_OPTIONS:
            raise ValueError(f"{row_label}: raw_policy must be one of {', '.join(POLICY_OPTIONS)}")
        latest_enabled = required_bool(
            sheet.cell(row=row_number, column=column_index["latest_enabled"]).value,
            "latest_enabled",
            row_label,
        )
        reason = str(sheet.cell(row=row_number, column=column_index["reason"]).value or "").strip()
        rows.append(
            SnapshotPolicyRow(
                point=point,
                raw_enabled=raw_enabled,
                raw_policy=raw_policy,
                latest_enabled=latest_enabled,
                reason=reason,
            )
        )

    return rows


def point_key(point: PlcPoint) -> tuple[int, str, str]:
    return (point.db_number, point.group_name, point.name)


def format_policy_sheet(
    sheet: Any,
    raw_enabled_validation: DataValidation,
    policy_validation: DataValidation,
    latest_enabled_validation: DataValidation,
) -> None:
    sheet.freeze_panes = "A2"
    for column_index, column_name in enumerate(POLICY_COLUMNS, start=1):
        width = 14
        if column_name in {"name", "source_name", "desc", "reason"}:
            width = 32
        if column_name == "group_name":
            width = 24
        sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = width

    if sheet.max_row >= 2:
        raw_enabled_column = column_letter_for(sheet, "raw_enabled")
        raw_policy_column = column_letter_for(sheet, "raw_policy")
        latest_enabled_column = column_letter_for(sheet, "latest_enabled")
        raw_enabled_validation.add(f"{raw_enabled_column}2:{raw_enabled_column}{sheet.max_row}")
        policy_validation.add(f"{raw_policy_column}2:{raw_policy_column}{sheet.max_row}")
        latest_enabled_validation.add(f"{latest_enabled_column}2:{latest_enabled_column}{sheet.max_row}")


def column_letter_for(sheet: Any, column_name: str) -> str:
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    return sheet.cell(row=1, column=headers.index(column_name) + 1).column_letter


def policy_sheets(workbook: Any) -> list[Any]:
    if "SNAPSHOT_POLICY" in workbook.sheetnames:
        return [workbook["SNAPSHOT_POLICY"]]
    return [workbook[name] for name in workbook.sheetnames if name != "README"]


def safe_sheet_name(db_number: int, group_name: str, used_sheet_names: set[str]) -> str:
    safe_group = re.sub(r"[\[\]\:\*\?\/\\]", "_", group_name).strip() or "Group"
    base = f"DB{db_number}_{safe_group}"[:31]
    candidate = base
    suffix = 1
    while candidate in used_sheet_names:
        suffix_text = f"_{suffix}"
        candidate = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    used_sheet_names.add(candidate)
    return candidate
