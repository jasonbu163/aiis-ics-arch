"""
文件路径: /tools/plc/point-mapping/src-python/workbook/init_workbook.py
功能描述: point-mapping XLSX 初始化辅助工具
主要功能:
    - 根据 TIA DB 源文件生成 DB配置 工作簿
    - 根据已填写 DB配置 和 DB 源文件生成待粘贴 TIA 点表的工作簿骨架
    - 保留 TIA 复制表头和 Static 行，避免伪造 offset / bit 数据
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from source_parser.tia import parse_db_source


TOOL_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DB_CONFIG_INPUT_PATH = TOOL_ROOT / "inputs" / "db_config.xlsx"
DEFAULT_DB_CONFIG_OUTPUT_PATH = TOOL_ROOT / "outputs" / "db_config.xlsx"
DEFAULT_WORKBOOK_OUTPUT_PATH = TOOL_ROOT / "outputs" / "DB.xlsx"

CONFIG_SHEET_NAME = "DB配置"
DB_CONFIG_COLUMNS = ["sheet_name", "db_name", "db_number", "group_name", "enabled", "备注"]
POINT_SHEET_COLUMNS = [
    "名称",
    "数据类型",
    "偏移量",
    "起始值",
    "保持",
    "从HMI/OPC UA/Web API可访问",
    "从HMI/OPC UA/Web API可写",
    "在HMI工程组态中可见",
    "设定值",
    "监控",
    "注释",
]

HEADER_FILL = "D9EAF7"
INVALID_SHEET_CHARS = set("[]:*?/\\")


@dataclass(frozen=True)
class DbConfigEntry:
    sheet_name: str
    db_name: str
    db_number: int | None
    group_name: str
    enabled: bool
    note: str = ""


def init_db_config_workbook(
    output_path: Path,
    db_source_path: Path | None = None,
    *,
    overwrite: bool = False,
) -> dict[str, Any]:
    """Create a workbook containing only the DB configuration sheet."""
    _ensure_can_write(output_path, overwrite)
    block_names = list(parse_db_source(db_source_path).keys()) if db_source_path else []

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = CONFIG_SHEET_NAME
    _write_header(sheet, DB_CONFIG_COLUMNS)

    if block_names:
        for block_name in block_names:
            sheet.append([
                block_name,
                block_name,
                "",
                block_name,
                True,
                "填写 db_number 后再运行 init-workbook；一个 DB 块对应一个 sheet。",
            ])
    else:
        sheet.append(["", "", "", "", True, "填写 sheet_name、db_name、db_number、group_name 和 enabled。"])

    _finish_sheet(sheet, DB_CONFIG_COLUMNS)
    _save_workbook(workbook, output_path)
    return {
        "output": str(output_path),
        "db_blocks": len(block_names),
        "config_rows": len(block_names) if block_names else 1,
    }


def init_point_workbook(
    db_config_path: Path,
    db_source_path: Path,
    output_path: Path,
    *,
    overwrite: bool = False,
) -> dict[str, Any]:
    """Create one empty TIA-copy point sheet per enabled DB config row."""
    _ensure_can_write(output_path, overwrite)
    db_blocks = parse_db_source(db_source_path)
    entries = _read_db_config_entries(db_config_path)
    enabled_entries = [entry for entry in entries if entry.enabled]
    _validate_entries(enabled_entries, db_blocks)

    workbook = Workbook()
    config_sheet = workbook.active
    config_sheet.title = CONFIG_SHEET_NAME
    _write_header(config_sheet, DB_CONFIG_COLUMNS)
    for entry in entries:
        config_sheet.append([
            entry.sheet_name,
            entry.db_name,
            entry.db_number,
            entry.group_name,
            entry.enabled,
            entry.note,
        ])
    _finish_sheet(config_sheet, DB_CONFIG_COLUMNS)

    for entry in enabled_entries:
        sheet = workbook.create_sheet(entry.sheet_name)
        _write_header(sheet, POINT_SHEET_COLUMNS)
        sheet.append(["Static", "", "", "", "", "", "", "", "", "", ""])
        _finish_sheet(sheet, POINT_SHEET_COLUMNS)

    _save_workbook(workbook, output_path)
    return {
        "output": str(output_path),
        "db_config": str(db_config_path),
        "db_source": str(db_source_path),
        "db_blocks": len(db_blocks),
        "config_rows": len(entries),
        "enabled_sheets": len(enabled_entries),
    }


def _read_db_config_entries(db_config_path: Path) -> list[DbConfigEntry]:
    if not db_config_path.is_file():
        raise FileNotFoundError(f"DB config workbook not found: {db_config_path}")
    workbook = load_workbook(db_config_path, read_only=True, data_only=True)
    if CONFIG_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"{db_config_path}: missing {CONFIG_SHEET_NAME} sheet")

    sheet = workbook[CONFIG_SHEET_NAME]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{db_config_path}: {CONFIG_SHEET_NAME} sheet is empty")
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    missing = [column for column in DB_CONFIG_COLUMNS[:5] if column not in headers]
    if missing:
        raise ValueError(f"{db_config_path}: {CONFIG_SHEET_NAME} missing columns: {', '.join(missing)}")
    column_index = {header: index for index, header in enumerate(headers)}

    entries: list[DbConfigEntry] = []
    for excel_row, row in enumerate(rows[1:], start=2):
        sheet_name = _cell_text(row, column_index["sheet_name"])
        if not sheet_name:
            continue
        db_name = _cell_text(row, column_index["db_name"]) or sheet_name
        group_name = _cell_text(row, column_index["group_name"]) or db_name
        enabled = _normalize_enabled(_cell_value(row, column_index["enabled"]))
        db_number_value = _cell_value(row, column_index["db_number"])
        if _is_blank(db_number_value) and enabled:
            raise ValueError(f"{CONFIG_SHEET_NAME}!{excel_row}: db_number is required for {sheet_name!r}")
        db_number = None if _is_blank(db_number_value) else int(float(str(db_number_value).strip()))
        entries.append(DbConfigEntry(
            sheet_name=sheet_name,
            db_name=db_name,
            db_number=db_number,
            group_name=group_name,
            enabled=enabled,
            note=_cell_text(row, column_index.get("备注", -1)),
        ))
    if not entries:
        raise ValueError(f"{db_config_path}: no DB config rows found")
    return entries


def _validate_entries(entries: list[DbConfigEntry], db_blocks: dict[str, Any]) -> None:
    if not entries:
        raise ValueError("no enabled DB config rows found")
    seen_sheet_names: set[str] = set()
    for entry in entries:
        _validate_sheet_title(entry.sheet_name)
        if entry.sheet_name in seen_sheet_names:
            raise ValueError(f"duplicate sheet_name in DB config: {entry.sheet_name}")
        seen_sheet_names.add(entry.sheet_name)
        if entry.db_name not in db_blocks:
            raise ValueError(f"{entry.sheet_name}: DATA_BLOCK not found for DB name {entry.db_name!r}")


def _validate_sheet_title(sheet_name: str) -> None:
    if len(sheet_name) > 31:
        raise ValueError(f"sheet_name exceeds Excel 31-character limit: {sheet_name!r}")
    if any(char in INVALID_SHEET_CHARS for char in sheet_name):
        raise ValueError(f"sheet_name contains invalid Excel characters: {sheet_name!r}")


def _ensure_can_write(output_path: Path, overwrite: bool) -> None:
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"output workbook already exists: {output_path}; pass --overwrite to replace it")


def _save_workbook(workbook: Workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _write_header(sheet: Any, headers: list[str]) -> None:
    sheet.append(headers)
    fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill


def _finish_sheet(sheet: Any, headers: list[str]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, header in enumerate(headers, start=1):
        width = 14
        if header in {"sheet_name", "db_name", "group_name", "名称", "数据类型"}:
            width = 24
        if header in {"备注", "注释", "从HMI/OPC UA/Web API可访问", "从HMI/OPC UA/Web API可写", "在HMI工程组态中可见"}:
            width = 32
        sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = width


def _cell_value(row: tuple[Any, ...], index: int) -> Any:
    if index < 0 or index >= len(row):
        return None
    return row[index]


def _cell_text(row: tuple[Any, ...], index: int) -> str:
    value = _cell_value(row, index)
    if _is_blank(value):
        return ""
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _normalize_enabled(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if _is_blank(value):
        return True
    text = str(value).strip().casefold()
    if text in {"true", "1", "yes"}:
        return True
    if text in {"false", "0", "no"}:
        return False
    return True
