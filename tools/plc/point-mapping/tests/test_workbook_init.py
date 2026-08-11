"""Sanitized tests for the point-mapping workbook scaffolding."""
from __future__ import annotations

from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook


TOOL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOL_ROOT / "src-python"))

from workbook.init_workbook import (  # noqa: E402
    CONFIG_SHEET_NAME,
    DB_CONFIG_COLUMNS,
    POINT_SHEET_COLUMNS,
    init_db_config_workbook,
    init_point_workbook,
)


DB_SOURCE = '''
DATA_BLOCK "Example_Main"
{ S7_Optimized_Access := 'FALSE' }
VERSION : 0.1
STRUCT
  Value : Real;
END_STRUCT;
END_DATA_BLOCK

DATA_BLOCK "Example_Aux"
{ S7_Optimized_Access := 'FALSE' }
VERSION : 0.1
STRUCT
  Status : Bool;
END_STRUCT;
END_DATA_BLOCK
'''


class WorkbookInitTests(unittest.TestCase):
    def test_init_db_config_prefills_sanitized_data_block_names(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "example.db"
            output_path = Path(temp_dir) / "db_config.xlsx"
            source_path.write_text(DB_SOURCE, encoding="utf-8")

            result = init_db_config_workbook(output_path, source_path)
            workbook = load_workbook(output_path, data_only=True)

        self.assertEqual(result["db_blocks"], 2)
        self.assertEqual(workbook.sheetnames, [CONFIG_SHEET_NAME])
        sheet = workbook[CONFIG_SHEET_NAME]
        self.assertEqual([cell.value for cell in sheet[1]], DB_CONFIG_COLUMNS)
        self.assertEqual(sheet["A2"].value, "Example_Main")
        self.assertEqual(sheet["B2"].value, "Example_Main")

    def test_init_workbook_creates_point_sheet_skeletons(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "example.db"
            db_config_path = root / "db_config.xlsx"
            output_path = root / "DB.xlsx"
            source_path.write_text(DB_SOURCE, encoding="utf-8")
            init_db_config_workbook(db_config_path, source_path)

            workbook = load_workbook(db_config_path)
            sheet = workbook[CONFIG_SHEET_NAME]
            headers = [cell.value for cell in sheet[1]]
            number_column = headers.index("db_number") + 1
            for row, number in enumerate((1, 2), start=2):
                sheet.cell(row=row, column=number_column).value = number
            workbook.save(db_config_path)

            result = init_point_workbook(db_config_path, source_path, output_path)
            generated = load_workbook(output_path, data_only=True)

        self.assertEqual(result["enabled_sheets"], 2)
        self.assertEqual(generated.sheetnames[0], CONFIG_SHEET_NAME)
        self.assertIn("Example_Main", generated.sheetnames)
        point_sheet = generated["Example_Main"]
        self.assertEqual([cell.value for cell in point_sheet[1]], POINT_SHEET_COLUMNS)
        self.assertEqual(point_sheet["A2"].value, "Static")

    def test_disabled_rows_are_not_materialized(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "example.db"
            db_config_path = root / "db_config.xlsx"
            output_path = root / "DB.xlsx"
            source_path.write_text(DB_SOURCE, encoding="utf-8")
            init_db_config_workbook(db_config_path, source_path)

            workbook = load_workbook(db_config_path)
            sheet = workbook[CONFIG_SHEET_NAME]
            headers = [cell.value for cell in sheet[1]]
            number_column = headers.index("db_number") + 1
            enabled_column = headers.index("enabled") + 1
            sheet.cell(row=2, column=number_column).value = 1
            sheet.cell(row=3, column=enabled_column).value = False
            workbook.save(db_config_path)

            result = init_point_workbook(db_config_path, source_path, output_path)
            generated = load_workbook(output_path, data_only=True)

        self.assertEqual(result["enabled_sheets"], 1)
        self.assertIn("Example_Main", generated.sheetnames)
        self.assertNotIn("Example_Aux", generated.sheetnames)


if __name__ == "__main__":
    unittest.main()
