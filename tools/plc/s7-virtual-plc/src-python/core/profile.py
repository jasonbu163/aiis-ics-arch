"""
文件路径: /tools/plc/s7-virtual-plc/src-python/core/profile.py
功能描述: simulation_profile.xlsx 模板生成和规则读取
"""
from __future__ import annotations

import math
import time
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core.models import (
    MODE_OPTIONS,
    NUMERIC_TYPES,
    PROFILE_COLUMNS,
    SequenceStep,
    SimulationRule,
    VirtualDbArea,
    VirtualPoint,
)


def number_from_name(text: str, default: int = 1) -> int:
    digits = "".join(char for char in text if char.isdigit())
    return int(digits[-2:]) if digits else default


def profile_key(point: VirtualPoint) -> str:
    return f"DB{point.db_number}.{point.name}"


def default_simulated_value(point: VirtualPoint, tick: int, started_at: float) -> Any:
    elapsed = time.time() - started_at
    name = point.name.lower()
    source = point.source_name.lower()
    text = f"{name} {source} {point.desc.lower()}"
    index = number_from_name(point.name, default=(point.offset % 9) + 1)
    wave = math.sin((tick + index) / 5.0)

    if point.data_type == "Bool":
        if any(token in text for token in ("fault", "fail", "ovl", "estop", "故障", "过载", "急停")):
            return False
        if any(token in text for token in ("running", "ready", "enable", "运行", "准备", "使能")):
            return True
        return (tick + point.offset + point.bit) % 17 == 0

    if point.data_type in {"String"}:
        if "coilnumber" in name or "coilnumber" in source or "卷材编号" in text:
            return f"MOCK{datetime.now().strftime('%H%M%S')}"
        if "coiltype" in name or "卷材种类" in text:
            return "MOCK-STEEL"
        return f"SIM-{point.db_number}-{point.offset}"

    if point.data_type == "Date_And_Time":
        return datetime.now()

    if "temset" in name or "设定温度" in text:
        value = 150 + index * 8
    elif "temact" in name or "实际温度" in text:
        value = 150 + index * 8 + wave * 3
    elif "output" in name:
        value = 45 + wave * 15
    elif "speed" in name or "速度" in text:
        value = 18 + wave * 2
    elif "tension" in name or "张力" in text:
        value = 55 + wave * 5
    elif "dia" in name or "卷径" in text:
        value = 420 + ((elapsed * 1.5 + index * 11) % 180)
    elif "length" in name or "长度" in text:
        value = elapsed * 1.2
    elif "width" in name or "宽度" in text:
        value = 1250
    elif "thickness" in name or "厚度" in text:
        value = 0.8
    elif "liquid" in name or "液位" in text:
        value = 320 + wave * 12
    else:
        value = (point.offset % 100) + wave

    if point.data_type in {"Int", "DInt"}:
        return int(round(value))
    if point.data_type == "USInt":
        return max(0, min(255, int(round(value))))
    if point.data_type in {"UInt", "Word"}:
        return max(0, min(65535, int(round(value))))
    if point.data_type == "UDInt":
        return max(0, min(4294967295, int(round(value))))
    return round(float(value), 2)


def is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def parse_bool(value: Any, *, default: bool | None = None) -> bool | None:
    if isinstance(value, bool):
        return value
    if is_blank(value):
        return default
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "y", "是", "启用"}:
        return True
    if text in {"false", "0", "no", "n", "否", "禁用"}:
        return False
    return default


def parse_float(value: Any, field_name: str, row_label: str) -> float:
    if is_blank(value):
        raise ValueError(f"{row_label}: {field_name} is required")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{row_label}: {field_name} must be numeric") from exc


def coerce_value(point: VirtualPoint, value: Any, row_label: str) -> Any:
    if point.data_type == "Bool":
        parsed = parse_bool(value)
        if parsed is None:
            raise ValueError(f"{row_label}: Bool value must be true/false")
        return parsed
    if point.data_type in NUMERIC_TYPES:
        number = parse_float(value, "value", row_label)
        if point.data_type == "Real":
            return round(number, 2)
        return int(round(number))
    if point.data_type == "Date_And_Time":
        if isinstance(value, datetime):
            return value
        if is_blank(value) or str(value).strip().lower() == "now":
            return datetime.now()
        try:
            return datetime.fromisoformat(str(value).strip())
        except ValueError as exc:
            raise ValueError(f"{row_label}: Date_And_Time value must be ISO datetime or now") from exc
    return "" if value is None else str(value)


def format_default_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return "now"
    return value


def default_profile_row(point: VirtualPoint) -> dict[str, Any]:
    value = default_simulated_value(point, tick=0, started_at=time.time())
    row: dict[str, Any] = {
        "enabled": True,
        "db_number": point.db_number,
        "group_name": point.group_name,
        "name": point.name,
        "source_name": point.source_name,
        "type": point.data_type,
        "offset": point.offset,
        "bit": point.bit if point.data_type == "Bool" else "",
        "desc": point.desc,
        "mode": "fixed",
        "min_value": "",
        "max_value": "",
        "fixed_value": format_default_value(value),
        "hold_seconds": "",
        "period_seconds": "",
        "sequence_values": "",
    }
    if point.data_type in NUMERIC_TYPES:
        base = float(value)
        span = max(1.0, abs(base) * 0.05)
        row.update({
            "mode": "random_range",
            "min_value": round(base - span, 2),
            "max_value": round(base + span, 2),
            "fixed_value": "",
            "hold_seconds": 5,
        })
    if point.data_type == "Date_And_Time":
        row.update({"mode": "now", "fixed_value": ""})
    return row


def write_profile_template(areas: dict[int, VirtualDbArea], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["S7 virtual PLC simulation profile"])
    readme.append(["1. This file is generated from tools/config/plc_points.yaml."])
    readme.append(["2. Edit the DB sheets, then copy this file to tools/plc/s7-virtual-plc/inputs/simulation_profile.xlsx."])
    readme.append(["3. Normal startup reads inputs/simulation_profile.xlsx and exits on validation errors."])
    readme.append(["4. Only fields required by the selected mode are read. Leave unrelated fields blank."])
    readme.append(["Modes", "Required fields"])
    readme.append(["fixed", "fixed_value"])
    readme.append(["random_range", "min_value, max_value, hold_seconds"])
    readme.append(["wave", "min_value, max_value, period_seconds"])
    readme.append(["sequence", "sequence_values as value:seconds;value:seconds"])
    readme.append(["now", "Date_And_Time points only"])
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 80

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for area in areas.values():
        sheet = workbook.create_sheet(title=f"DB{area.db_number}")
        mode_validation = DataValidation(type="list", formula1=f'"{",".join(MODE_OPTIONS)}"', allow_blank=False)
        enabled_validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
        sheet.add_data_validation(mode_validation)
        sheet.add_data_validation(enabled_validation)
        sheet.append(PROFILE_COLUMNS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for point in area.points:
            row = default_profile_row(point)
            sheet.append([row[column] for column in PROFILE_COLUMNS])
        sheet.freeze_panes = "A2"
        for column_index, column_name in enumerate(PROFILE_COLUMNS, start=1):
            width = 14
            if column_name in {"name", "source_name", "desc", "sequence_values"}:
                width = 26
            sheet.column_dimensions[sheet.cell(row=1, column=column_index).column_letter].width = width
        if sheet.max_row >= 2:
            enabled_validation.add(f"A2:A{sheet.max_row}")
            mode_validation.add(f"J2:J{sheet.max_row}")

    workbook.save(output_path)


def read_profile_sheet(sheet: Any, expected_points: dict[str, VirtualPoint]) -> dict[str, SimulationRule]:
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    missing = [column for column in PROFILE_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"{sheet.title}: missing columns: {', '.join(missing)}")
    column_index = {name: headers.index(name) + 1 for name in PROFILE_COLUMNS}
    rules: dict[str, SimulationRule] = {}
    seen_points: set[str] = set()

    for row_number in range(2, sheet.max_row + 1):
        name = sheet.cell(row=row_number, column=column_index["name"]).value
        if is_blank(name):
            continue
        point_name = str(name).strip()
        point = expected_points.get(point_name)
        row_label = f"{sheet.title} row {row_number} ({point_name})"
        if point is None:
            raise ValueError(f"{row_label}: point is not defined in plc_points.yaml")

        db_number = int(parse_float(sheet.cell(row=row_number, column=column_index["db_number"]).value, "db_number", row_label))
        point_type = str(sheet.cell(row=row_number, column=column_index["type"]).value or "").strip()
        offset = int(parse_float(sheet.cell(row=row_number, column=column_index["offset"]).value, "offset", row_label))
        if db_number != point.db_number or point_type != point.data_type or offset != point.offset:
            raise ValueError(f"{row_label}: db_number/type/offset does not match plc_points.yaml")

        enabled = parse_bool(sheet.cell(row=row_number, column=column_index["enabled"]).value, default=True)
        mode = str(sheet.cell(row=row_number, column=column_index["mode"]).value or "").strip()
        if mode not in MODE_OPTIONS:
            raise ValueError(f"{row_label}: mode must be one of {', '.join(MODE_OPTIONS)}")

        rule = SimulationRule(enabled=bool(enabled), mode=mode)

        if mode == "fixed":
            rule.fixed_value = coerce_value(point, sheet.cell(row=row_number, column=column_index["fixed_value"]).value, row_label)
        elif mode == "random_range":
            if point.data_type not in NUMERIC_TYPES:
                raise ValueError(f"{row_label}: random_range only supports numeric point types")
            rule.hold_seconds = max(
                0.1,
                parse_float(sheet.cell(row=row_number, column=column_index["hold_seconds"]).value, "hold_seconds", row_label),
            )
            rule.min_value = parse_float(sheet.cell(row=row_number, column=column_index["min_value"]).value, "min_value", row_label)
            rule.max_value = parse_float(sheet.cell(row=row_number, column=column_index["max_value"]).value, "max_value", row_label)
            if rule.min_value > rule.max_value:
                raise ValueError(f"{row_label}: min_value must be <= max_value")
        elif mode == "wave":
            if point.data_type not in NUMERIC_TYPES:
                raise ValueError(f"{row_label}: wave only supports numeric point types")
            rule.period_seconds = max(
                0.1,
                parse_float(sheet.cell(row=row_number, column=column_index["period_seconds"]).value, "period_seconds", row_label),
            )
            rule.min_value = parse_float(sheet.cell(row=row_number, column=column_index["min_value"]).value, "min_value", row_label)
            rule.max_value = parse_float(sheet.cell(row=row_number, column=column_index["max_value"]).value, "max_value", row_label)
            if rule.min_value > rule.max_value:
                raise ValueError(f"{row_label}: min_value must be <= max_value")
        elif mode == "sequence":
            hold_value = sheet.cell(row=row_number, column=column_index["hold_seconds"]).value
            if not is_blank(hold_value):
                rule.hold_seconds = max(0.1, parse_float(hold_value, "hold_seconds", row_label))
            sequence_text = str(sheet.cell(row=row_number, column=column_index["sequence_values"]).value or "").strip()
            if not sequence_text:
                raise ValueError(f"{row_label}: sequence_values is required for sequence mode")
            rule.sequence_values = parse_sequence_values(point, sequence_text, row_label, rule.hold_seconds)
        elif mode == "now":
            if point.data_type != "Date_And_Time":
                raise ValueError(f"{row_label}: now only supports Date_And_Time point types")

        rules[profile_key(point)] = rule
        seen_points.add(point_name)

    missing_points = sorted(set(expected_points) - seen_points)
    if missing_points:
        raise ValueError(f"{sheet.title}: missing points from XLSX: {', '.join(missing_points[:5])}")
    return rules


def parse_sequence_values(point: VirtualPoint, text: str, row_label: str, default_duration: float) -> list[SequenceStep]:
    steps: list[SequenceStep] = []
    for raw_step in text.split(";"):
        step = raw_step.strip()
        if not step:
            continue
        if ":" in step:
            raw_value, raw_duration = step.rsplit(":", 1)
            duration = parse_float(raw_duration.strip(), "sequence duration", row_label)
        else:
            raw_value = step
            duration = default_duration
        if duration <= 0:
            raise ValueError(f"{row_label}: sequence duration must be > 0")
        steps.append(SequenceStep(value=coerce_value(point, raw_value.strip(), row_label), duration_seconds=duration))
    if not steps:
        raise ValueError(f"{row_label}: sequence_values is empty")
    return steps


def load_profile_rules(areas: dict[int, VirtualDbArea], profile_path: Path) -> dict[str, SimulationRule]:
    if not profile_path.exists():
        raise FileNotFoundError(
            f"Simulation profile not found: {profile_path}. "
            "Run init-profile, edit the generated XLSX, then copy it into inputs/."
        )

    workbook = load_workbook(profile_path, data_only=True)
    rules: dict[str, SimulationRule] = {}
    for area in areas.values():
        sheet_name = f"DB{area.db_number}"
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Simulation profile is missing sheet: {sheet_name}")
        expected_points = {point.name: point for point in area.points}
        rules.update(read_profile_sheet(workbook[sheet_name], expected_points))
    return rules
